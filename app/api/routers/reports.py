import datetime
import io
from typing import Optional
from fastapi import APIRouter, Depends
from fastapi.responses import Response
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from app.db.database import get_db
from app.api.deps import get_current_user, require_admin
from app.crud.crud_ticket import get_tickets
from app.crud.crud_category import get_category_map
from app.crud.crud_team import get_team_map
from app.crud.crud_settings import get_sla_config
from app.services.formatting import enrich
from app.services.reporting import build_report_data

router = APIRouter(prefix="/reports", tags=["reports"])

_XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

def _autosize(ws):
    for col in ws.columns:
        width = max((len(str(c.value)) if c.value is not None else 0) for c in col) + 2
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width, 10), 50)

def _xlsx_response(wb: Workbook, filename: str) -> Response:
    buf = io.BytesIO()
    wb.save(buf)
    return Response(content=buf.getvalue(), media_type=_XLSX_MEDIA,
                     headers={"Content-Disposition": f'attachment; filename="{filename}"'})

def _write_rows(ws, headers, rows):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(r)

@router.get("/tickets.xlsx")
def export_tickets(status: str = "", team: str = "", scope: str = "", q_: str = "",
                    priority: str = "", category: str = "", mmu_vehicle: str = "", district: str = "",
                    district_id: Optional[int] = None, mandal_id: Optional[int] = None,
                    date_from: str = "", date_to: str = "", time_from: str = "", time_to: str = "", shift: str = "",
                    sort_by: str = "", sort_desc: bool = False,
                    db: Session = Depends(get_db), current_user: dict = Depends(get_current_user)):
    """Exports the same rows the Ticket Queue would show for these filters -
    reuses get_tickets, so a plain team member only ever exports their own
    team's data, exactly like the queue itself already scopes access."""
    tickets, total = get_tickets(db, current_user, status=status, team=team, scope=scope, q=q_,
                                  priority=priority, category=category, mmu_vehicle=mmu_vehicle, district=district,
                                  district_id=district_id, mandal_id=mandal_id,
                                  date_from=date_from, date_to=date_to, time_from=time_from, time_to=time_to, shift=shift,
                                  page=1, page_size=5000, sort_by=sort_by, sort_desc=sort_desc)
    category_map, team_map, sla_cfg = get_category_map(db), get_team_map(db), get_sla_config(db)
    rows = [enrich(t, category_map, team_map, sla_cfg) for t in tickets]

    wb = Workbook()
    ws = wb.active
    ws.title = "Tickets"
    headers = ["Ticket No", "Source", "MMU/Vehicle", "District", "Category", "Priority", "Team", "Status",
               "Created", "Due (TAT)", "Closed", "TAT (min)", "Breached", "Escalated", "Assignee", "Resolution"]
    _write_rows(ws, headers, [
        [r.get("ticket_no"), r.get("source"), r.get("mmu_vehicle"), r.get("district"), r.get("category_label"),
         r.get("priority"), r.get("team_label"), r.get("status"), r.get("created_at"), r.get("due_at"),
         r.get("closed_at"), r.get("tat_mins"), bool(r.get("breached")), bool(r.get("escalated")),
         r.get("assignee"), r.get("resolution")]
        for r in rows
    ])
    _autosize(ws)
    return _xlsx_response(wb, f"tickets_{datetime.date.today().isoformat()}.xlsx")

def _period_range(period: str):
    today = datetime.date.today()
    if period == "weekly":
        return (today - datetime.timedelta(days=6)).isoformat(), today.isoformat()
    if period == "monthly":
        return (today - datetime.timedelta(days=29)).isoformat(), today.isoformat()
    return today.isoformat(), today.isoformat()

@router.get("/summary.xlsx")
def export_summary(period: str = "daily", district: str = "", mmu_vehicle: str = "", team: str = "",
                    date_from: str = "", date_to: str = "",
                    db: Session = Depends(get_db), current_user: dict = Depends(require_admin)):
    """Management report: the same figures the Daily Monitoring dashboard
    shows, one section per sheet. `period` is only a convenience that fills
    in date_from/date_to when they aren't given explicitly."""
    if not date_from and not date_to:
        date_from, date_to = _period_range(period)
    d = build_report_data(db, district=district, mmu_vehicle=mmu_vehicle, team=team,
                           date_from=date_from, date_to=date_to)

    wb = Workbook()
    ws = wb.active
    ws.title = "Overview"
    ws.append(["Report period", f"{date_from} to {date_to}"])
    ws.append(["District filter", district or "(all)"])
    ws.append(["MMU/Vehicle filter", mmu_vehicle or "(all)"])
    ws.append(["Team filter", team or "(all)"])
    ws.append([])
    ws.append(["Metric", "Value"])
    for cell in ws[6]:
        cell.font = Font(bold=True)
    for k, v in d["today"].items():
        ws.append([k.replace("_", " ").title(), v])
    ws.append([])
    ws.append(["KPI", "Value"])
    for k, v in d["kpi"].items():
        ws.append([k.replace("_", " ").title(), v])
    _autosize(ws)

    ws2 = wb.create_sheet("By Category")
    _write_rows(ws2, ["Category", "Total", "Open"], [[r["category"], r["n"], r["open_n"]] for r in d["by_category"]])
    _autosize(ws2)

    ws3 = wb.create_sheet("By Team")
    _write_rows(ws3, ["Team", "Total", "Open", "Breached", "Unassigned"],
                [[r["team"], r["n"], r["open_n"], r["breach_n"], r["unassigned_n"]] for r in d["by_team"]])
    _autosize(ws3)

    ws4 = wb.create_sheet("By Priority")
    _write_rows(ws4, ["Priority", "Total", "Open"], [[r["priority"], r["n"], r["open_n"]] for r in d["by_priority"]])
    _autosize(ws4)

    ws5 = wb.create_sheet("By Status")
    _write_rows(ws5, ["Status", "Count"], [[r["status"], r["n"]] for r in d["by_status"]])
    _autosize(ws5)

    ws6 = wb.create_sheet("Category x MMU")
    _write_rows(ws6, ["Category", "MMU/Vehicle", "Tickets"],
                [[r["category"], r["mmu_vehicle"], r["n"]] for r in d["by_category_mmu"]])
    _autosize(ws6)

    ws7 = wb.create_sheet("Repeat Vehicles")
    _write_rows(ws7, ["MMU/Vehicle", "Tickets"], [[r["mmu_vehicle"], r["n"]] for r in d["repeat_vehicles"]])
    _autosize(ws7)

    ws8 = wb.create_sheet("Chronic Equipment")
    _write_rows(ws8, ["MMU/Vehicle", "District", "Breakdowns (30d)", "Still Open", "Issue Types", "Last Breakdown"],
                [[r["mmu_vehicle"], r["district"], r["n"], r["open_n"], r["categories"], r["last_breakdown_at"]]
                 for r in d["chronic_equipment"]])
    _autosize(ws8)

    return _xlsx_response(wb, f"summary_{period}_{date_from}_to_{date_to}.xlsx")
