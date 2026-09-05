import io
from openpyxl import load_workbook
from tests.conftest import auth_headers

CT = auth_headers("CALL_TAKER")
SVC = auth_headers("SERVICE")
MGR = auth_headers("CC_MANAGER")


def make_ticket(client, mmu_vehicle="AP1REPORT", district="Report District", category="MACHINE", priority="P2"):
    r = client.post("/cccapi/ticket", json={
        "mmu_vehicle": mmu_vehicle, "district": district, "problem": "reporting test",
        "category": category, "priority": priority, "caller_name": "Test Caller", "caller_phone": "9876543210",
    }, headers=CT)
    assert r.status_code == 200, r.text
    return r.json()["id"]


def test_dashboard_with_no_filters_still_works(client):
    make_ticket(client)
    r = client.get("/cccapi/dashboard", headers=MGR)
    assert r.status_code == 200
    body = r.json()
    assert "today" in body and "by_category" in body and "kpi" in body
    assert "by_category_mmu" in body  # new breakdown


def test_dashboard_district_filter_narrows_results(client):
    make_ticket(client, mmu_vehicle="AP1NARROW", district="Narrow District Alpha")
    make_ticket(client, mmu_vehicle="AP2NARROW", district="Narrow District Beta")

    r_all = client.get("/cccapi/dashboard", headers=MGR).json()
    r_alpha = client.get("/cccapi/dashboard?district=Narrow District Alpha", headers=MGR).json()
    assert r_alpha["today"]["tickets"] <= r_all["today"]["tickets"]
    assert r_alpha["today"]["tickets"] >= 1


def test_dashboard_mmu_vehicle_filter(client):
    tid = make_ticket(client, mmu_vehicle="AP9UNIQUE9")
    r = client.get("/cccapi/dashboard?mmu_vehicle=AP9UNIQUE9", headers=MGR).json()
    assert r["today"]["tickets"] >= 1
    found = [row for row in r["by_category_mmu"] if row["mmu_vehicle"] == "AP9UNIQUE9"]
    assert found and found[0]["n"] >= 1


def test_dashboard_team_filter(client):
    make_ticket(client, category="MACHINE")  # routes to SERVICE
    r_service = client.get("/cccapi/dashboard?team=SERVICE", headers=MGR).json()
    r_application = client.get("/cccapi/dashboard?team=APPLICATION", headers=MGR).json()
    service_rows = [row for row in r_service["by_team"] if row["team"] == "SERVICE"]
    application_rows = [row for row in r_application["by_team"] if row["team"] == "SERVICE"]
    assert service_rows  # SERVICE present when filtering by SERVICE
    assert not application_rows  # absent when filtering by a different team


def test_dashboard_forbidden_for_engineer(client):
    r = client.get("/cccapi/dashboard", headers=SVC)
    assert r.status_code == 403


def test_export_tickets_xlsx_is_a_valid_workbook(client):
    make_ticket(client)
    r = client.get("/cccapi/reports/tickets.xlsx", headers=MGR)
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Tickets"]
    assert wb["Tickets"]["A1"].value == "Ticket No"


def test_export_tickets_xlsx_engineer_only_sees_own_team(client):
    make_ticket(client, category="MACHINE")  # -> SERVICE
    r = client.get("/cccapi/reports/tickets.xlsx", headers=SVC)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Tickets"]
    teams_col = [ws.cell(row=i, column=7).value for i in range(2, ws.max_row + 1)]
    assert all(t == "Service Team" or t is None for t in teams_col)


def test_export_summary_xlsx_has_expected_sheets(client):
    make_ticket(client)
    r = client.get("/cccapi/reports/summary.xlsx?period=daily", headers=MGR)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Overview", "By Category", "By Team", "By Priority", "By Status", "Category x MMU", "Repeat Vehicles", "Chronic Equipment"]


def test_export_summary_xlsx_forbidden_for_engineer(client):
    r = client.get("/cccapi/reports/summary.xlsx", headers=SVC)
    assert r.status_code == 403
